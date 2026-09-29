from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..config import APP_NAME, APP_VERSION, REPORTS
from ..utils import new_id, p_format


STYLES = {
    "academic": {
        "title_prefix": "",
        "audience": "academic research",
        "tone": "Methods and results are reported in a scientific register. Causal language is avoided unless the design supports it.",
    },
    "business": {
        "title_prefix": "Analytics brief: ",
        "audience": "business analytics",
        "tone": "Findings are framed for decision support. Statistical caveats remain explicit.",
    },
    "scientific": {
        "title_prefix": "",
        "audience": "scientific analysis",
        "tone": "Emphasis on assumptions, uncertainty, and reproducibility.",
    },
    "general": {
        "title_prefix": "",
        "audience": "general analysis",
        "tone": "Plain-language interpretation sits beside the numerical results.",
    },
}


def build_report_payload(
    *,
    dataset_name: str,
    profile: dict,
    quality: dict,
    cleaning: list,
    question: str,
    recommendation: dict | None,
    assumptions: list,
    analysis: dict | None,
    validation: dict | None,
    interpretation: str,
    style: str = "academic",
) -> dict:
    meta = STYLES.get(style, STYLES["academic"])
    limitations = [
        "Results describe this sample under the stated model; they do not automatically generalize.",
        "Statistical significance is not practical significance.",
        "Association is not causation.",
    ]
    if validation:
        for f in validation.get("flags") or []:
            if f.get("level") in ("warning", "error"):
                limitations.append(f["message"])
    return {
        "title": meta["title_prefix"] + (question or f"Analysis of {dataset_name}"),
        "style": style,
        "audience": meta["audience"],
        "tone": meta["tone"],
        "generated_at": datetime.utcnow().isoformat() + "Z",
        "software": f"{APP_NAME} {APP_VERSION}",
        "dataset_name": dataset_name,
        "dataset_description": profile,
        "quality": quality,
        "cleaning": cleaning,
        "research_question": question,
        "recommendation": recommendation,
        "assumptions": assumptions,
        "analysis": analysis,
        "validation": validation,
        "interpretation": interpretation,
        "limitations": limitations,
        "conclusion": _conclusion(question, analysis, interpretation),
    }


def _conclusion(question, analysis, interpretation):
    if not analysis:
        return "No confirmatory analysis was executed in this report."
    return (
        (interpretation or analysis.get("interpretation") or analysis.get("summary") or "")
        + " See methods, assumptions, and validation flags before acting on this conclusion."
    )


def to_markdown(payload: dict) -> str:
    a = payload.get("analysis") or {}
    rec = payload.get("recommendation") or {}
    lines = [
        f"# {payload['title']}",
        "",
        f"*Generated {payload['generated_at']} by {payload['software']} · style: {payload['style']}*",
        "",
        "## 1. Dataset description",
        f"- File: **{payload['dataset_name']}**",
        f"- Rows: {payload['dataset_description'].get('n_rows')} · Columns: {payload['dataset_description'].get('n_cols')}",
        f"- Missing cells: {payload['dataset_description'].get('missing_pct')}%",
        f"- Duplicate rows: {payload['dataset_description'].get('duplicate_rows')}",
        f"- Quality score: {payload['dataset_description'].get('quality_score')}",
        "",
        "## 2. Data-quality assessment",
        f"- Issues: {payload.get('quality', {}).get('counts', {})}",
    ]
    for iss in (payload.get("quality") or {}).get("issues", [])[:20]:
        lines.append(f"- **{iss.get('severity')}** — {iss.get('problem')}")
    lines += ["", "## 3. Cleaning procedures"]
    if payload.get("cleaning"):
        for c in payload["cleaning"]:
            lines.append(f"- {c}")
    else:
        lines.append("- No cleaning operations were applied (or none were recorded).")
    lines += [
        "",
        "## 4. Research question",
        payload.get("research_question") or "_Not specified._",
        "",
        "## 5. Methodology",
        f"Recommended method: **{rec.get('recommended_label') or a.get('method_label') or a.get('method')}**",
        "",
    ]
    for w in rec.get("why") or []:
        lines.append(f"- {w}")
    if a.get("formula"):
        lines += ["", f"Formula: `{a['formula']}`"]
    lines += ["", "## 6. Statistical assumptions"]
    for c in payload.get("assumptions") or []:
        lines.append(f"- **{c.get('name')}** — {c.get('status')}: {c.get('evidence')}")
        lines.append(f"  - Action: {c.get('recommended_action')}")
    lines += ["", "## 7. Statistical analysis", a.get("summary") or "_No analysis._", ""]
    stats = a.get("statistics") or {}
    if stats:
        lines.append("### Key statistics")
        for k, v in stats.items():
            lines.append(f"- `{k}`: {v}")
    for tname, rows in (a.get("tables") or {}).items():
        if not rows:
            continue
        lines += ["", f"### Table: {tname}", ""]
        if isinstance(rows, list) and rows and isinstance(rows[0], dict):
            keys = list(rows[0].keys())
            lines.append("| " + " | ".join(keys) + " |")
            lines.append("| " + " | ".join("---" for _ in keys) + " |")
            for r in rows[:40]:
                lines.append("| " + " | ".join(str(r.get(k, ""))[:80] for k in keys) + " |")
    lines += [
        "",
        "## 8. Results",
        a.get("summary") or "",
        "",
        "## 9. Visualizations",
        "_Charts are available in the interactive workspace and can be exported separately._",
        "",
        "## 10. Interpretation",
        payload.get("interpretation") or a.get("interpretation") or "",
        "",
        "## 11. Limitations",
    ]
    for L in payload.get("limitations") or []:
        lines.append(f"- {L}")
    lines += ["", "## 12. Conclusion", payload.get("conclusion") or "", ""]
    if payload.get("validation"):
        lines += ["## Validation flags"]
        for f in payload["validation"].get("flags") or []:
            lines.append(f"- **{f.get('level')}**: {f.get('message')}")
    lines += [
        "",
        "---",
        "Every number in this report is produced by the analysis engine. "
        "The assistant does not invent p-values, coefficients, or citations.",
    ]
    return "\n".join(lines)


def to_html(payload: dict) -> str:
    md = to_markdown(payload)
    # lightweight
    import html as htmlmod

    body = htmlmod.escape(md)
    body = body.replace("\n", "<br>\n")
    return f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8"/>
<title>{htmlmod.escape(payload.get('title') or 'Report')}</title>
<style>
body {{ font-family: "Source Serif 4", Georgia, serif; max-width: 860px; margin: 40px auto; color: #1b2a32; background: #fbfaf7; line-height: 1.55; padding: 0 24px 80px; }}
h1 {{ font-family: "IBM Plex Sans", sans-serif; color: #0F4C5C; }}
code {{ background: #eef3f4; padding: 1px 4px; }}
</style></head>
<body>
<pre style="white-space:pre-wrap;font-family:inherit">{body}</pre>
</body></html>"""


def write_markdown(payload: dict, path: Path):
    path.write_text(to_markdown(payload), encoding="utf-8")
    return path


def write_html(payload: dict, path: Path):
    path.write_text(to_html(payload), encoding="utf-8")
    return path


def write_docx(payload: dict, path: Path):
    doc = Document()
    styles = doc.styles["Normal"]
    styles.font.name = "Calibri"
    styles.font.size = Pt(11)
    doc.add_heading(payload.get("title") or "Analysis report", 0)
    p = doc.add_paragraph(f"Generated {payload['generated_at']} by {payload['software']}")
    p.runs[0].italic = True
    doc.add_heading("Research question", 1)
    doc.add_paragraph(payload.get("research_question") or "")
    doc.add_heading("Methodology", 1)
    rec = payload.get("recommendation") or {}
    a = payload.get("analysis") or {}
    doc.add_paragraph(rec.get("recommended_label") or a.get("method_label") or "")
    for w in rec.get("why") or []:
        doc.add_paragraph(w, style="List Bullet")
    doc.add_heading("Assumptions", 1)
    for c in payload.get("assumptions") or []:
        doc.add_paragraph(f"{c.get('name')} [{c.get('status')}] — {c.get('evidence')}", style="List Bullet")
    doc.add_heading("Results", 1)
    doc.add_paragraph(a.get("summary") or "")
    stats = a.get("statistics") or {}
    if stats:
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Statistic"
        table.rows[0].cells[1].text = "Value"
        for k, v in list(stats.items())[:40]:
            row = table.add_row().cells
            row[0].text = str(k)
            row[1].text = str(v)
    doc.add_heading("Interpretation", 1)
    doc.add_paragraph(payload.get("interpretation") or a.get("interpretation") or "")
    doc.add_heading("Limitations", 1)
    for L in payload.get("limitations") or []:
        doc.add_paragraph(L, style="List Bullet")
    doc.add_heading("Conclusion", 1)
    doc.add_paragraph(payload.get("conclusion") or "")
    doc.save(str(path))
    return path


def write_pdf(payload: dict, path: Path):
    styles = getSampleStyleSheet()
    title = ParagraphStyle("T", parent=styles["Title"], textColor=colors.HexColor("#0F4C5C"), fontSize=16, leading=20)
    h = ParagraphStyle("H", parent=styles["Heading2"], textColor=colors.HexColor("#0F4C5C"), fontSize=12, spaceBefore=10)
    body = ParagraphStyle("B", parent=styles["Normal"], fontSize=9, leading=13)
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=0.7 * inch, rightMargin=0.7 * inch, topMargin=0.6 * inch, bottomMargin=0.6 * inch)
    story = [
        Paragraph(payload.get("title") or "Report", title),
        Paragraph(f"{payload['software']} · {payload['generated_at']}", body),
        Spacer(1, 8),
        Paragraph("Research question", h),
        Paragraph(payload.get("research_question") or "—", body),
        Paragraph("Results", h),
        Paragraph((payload.get("analysis") or {}).get("summary") or "—", body),
        Paragraph("Interpretation", h),
        Paragraph(payload.get("interpretation") or (payload.get("analysis") or {}).get("interpretation") or "—", body),
        Paragraph("Limitations", h),
    ]
    for L in payload.get("limitations") or []:
        story.append(Paragraph(f"• {L}", body))
    story += [Paragraph("Conclusion", h), Paragraph(payload.get("conclusion") or "—", body)]
    doc.build(story)
    return path


def write_xlsx(payload: dict, path: Path):
    a = payload.get("analysis") or {}
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame([{"key": k, "value": str(v)} for k, v in (a.get("statistics") or {}).items()]).to_excel(
            xw, sheet_name="statistics", index=False
        )
        pd.DataFrame(payload.get("assumptions") or []).to_excel(xw, sheet_name="assumptions", index=False)
        for name, rows in (a.get("tables") or {}).items():
            if isinstance(rows, list) and rows:
                pd.DataFrame(rows).to_excel(xw, sheet_name=str(name)[:28], index=False)
    return path


def export_report(payload: dict, fmt: str) -> Path:
    rid = new_id()[:10]
    REPORTS.mkdir(parents=True, exist_ok=True)
    if fmt == "md" or fmt == "markdown":
        p = REPORTS / f"report_{rid}.md"
        return write_markdown(payload, p)
    if fmt == "html":
        p = REPORTS / f"report_{rid}.html"
        return write_html(payload, p)
    if fmt == "docx":
        p = REPORTS / f"report_{rid}.docx"
        return write_docx(payload, p)
    if fmt == "pdf":
        p = REPORTS / f"report_{rid}.pdf"
        return write_pdf(payload, p)
    if fmt in ("xlsx", "excel"):
        p = REPORTS / f"report_{rid}.xlsx"
        return write_xlsx(payload, p)
    if fmt == "csv":
        p = REPORTS / f"report_{rid}.csv"
        stats = payload.get("analysis", {}).get("statistics") or {}
        pd.DataFrame([stats]).to_csv(p, index=False)
        return p
    if fmt == "json":
        p = REPORTS / f"report_{rid}.json"
        p.write_text(json.dumps(payload, default=str, indent=2), encoding="utf-8")
        return p
    raise ValueError(f"Unsupported format {fmt}")
